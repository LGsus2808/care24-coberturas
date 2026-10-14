import os
import re
import time
import json
from pathlib import Path
import docx
import pypdf
from groq import Groq

CARPETA_DOCUMENTOS = r"C:\Users\User\OneDrive - Care24 SA\CUENTAS CABINA MÉDICA (1)\CUENTAS CABINA MÉDICA"

GROQ_API_KEY = os.getenv("GROQ_API_KEY")
client = Groq(api_key=GROQ_API_KEY)

MODELO = "qwen/qwen3.8-27b"

def leer_word_completo(ruta):
    try:
        doc = docx.Document(ruta)
        lineas = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for tabla in doc.tables:
            for fila in tabla.rows:
                fila_txt = " | ".join(c.text.strip() for c in fila.cells if c.text.strip())
                if fila_txt:
                    lineas.append(fila_txt)
        return "\n".join(lineas)
    except Exception as e:
        print(f"⚠️ Error en Word {ruta.name}: {e}")
        return ""

def leer_pdf_completo(ruta):
    try:
        lector = pypdf.PdfReader(ruta)
        lineas = []
        for pag in lector.pages:
            try:
                txt = pag.extract_text()
                if txt:
                    lineas.append(txt)
            except Exception:
                pass
        return "\n".join(lineas)
    except Exception as e:
        print(f"⚠️ Error en PDF {ruta.name}: {e}")
        return ""

def dividir_en_bloques(texto, tamano_bloque=8000):
    return [texto[i:i + tamano_bloque] for i in range(0, len(texto), tamano_bloque)]

PROMPT_SISTEMA = """
Eres un especialista en pólizas de seguros y asistencias.
Tu tarea es extraer todas las coberturas y responder ÚNICAMENTE con un objeto JSON válido sin texto adicional.

Formato:
{
  "nombre_seguro": "Nombre del seguro o null",
  "empresa": "Aseguradora o null",
  "beneficiarios": ["Personas cubiertas"],
  "alcance_geografico": "Zonas geográficas",
  "coberturas": [
    {
      "categoria": "Médica | Hogar | Vial | Legal | Escolar | etc.",
      "servicio": "Nombre del servicio",
      "copago": "Copago o coaseguro",
      "limite_monto": "Monto máximo",
      "eventos_anuales": "Tope anual o Sin límite",
      "lugares_o_red": "Clínicas, talleres, domicilio o red",
      "exclusiones": "Exclusiones clave"
    }
  ]
}
"""

def extraer_json(texto):
    match = re.search(r'\{.*\}', texto, re.DOTALL)
    if match:
        try:
            return json.loads(match.group(0))
        except Exception:
            pass
    try:
        return json.loads(texto.strip())
    except Exception:
        return None

def consultar_llm(bloque_texto):
    for intento in range(4):
        try:
            chat_completion = client.chat.completions.create(
                messages=[
                    {"role": "system", "content": PROMPT_SISTEMA},
                    {"role": "user", "content": f"Extrae el JSON de coberturas de este texto:\n\n{bloque_texto}"}
                ],
                model=MODELO,
                max_tokens=750,
                temperature=0.1
            )
            raw_text = chat_completion.choices[0].message.content
            data = extraer_json(raw_text)
            if data:
                return data
            else:
                time.sleep(2)
        except Exception as e:
            espera = (intento + 1) * 15
            print(f"⏳ Pausa por cuota OTPM ({espera}s)...")
            time.sleep(espera)
    return None

def procesar_archivos():
    ruta_base = Path(CARPETA_DOCUMENTOS)
    if not ruta_base.exists():
        print(f"❌ La ruta no existe: {CARPETA_DOCUMENTOS}")
        return

    datos_consolidados = []
    print(f"🔍 Escaneando documentos en: {CARPETA_DOCUMENTOS} ...\n")

    for archivo in ruta_base.rglob("*"):
        nom = archivo.name.lower()
        if ("condicionado" in nom or "condiciones" in nom) and archivo.suffix.lower() in [".docx", ".pdf"]:
            print(f"\n📄 Leyendo: {archivo.name}")
            
            texto = leer_word_completo(archivo) if archivo.suffix.lower() == ".docx" else leer_pdf_completo(archivo)
            if not texto.strip():
                print(f"⚠️ Archivo sin texto legible: {archivo.name}")
                continue

            bloques = dividir_en_bloques(texto, tamano_bloque=8000)
            print(f"🧩 Fragmentos: {len(bloques)}")

            seguro_consolidado = {
                "nombre_seguro": archivo.stem,
                "empresa": "No especificada",
                "beneficiarios": [],
                "alcance_geografico": "No especificado",
                "coberturas": [],
                "archivo_fuente": archivo.name
            }

            for idx, bloque in enumerate(bloques, start=1):
                print(f"   -> Fragmento {idx}/{len(bloques)}...")
                data = consultar_llm(bloque)
                if data:
                    if data.get("nombre_seguro"): 
                        seguro_consolidado["nombre_seguro"] = data["nombre_seguro"]
                    if data.get("empresa"): 
                        seguro_consolidado["empresa"] = data["empresa"]
                    if data.get("alcance_geografico") and data["alcance_geografico"] != "No especificado":
                        seguro_consolidado["alcance_geografico"] = data["alcance_geografico"]
                    
                    for b in data.get("beneficiarios", []):
                        if b and b not in seguro_consolidado["beneficiarios"]:
                            seguro_consolidado["beneficiarios"].append(b)

                    for c in data.get("coberturas", []):
                        seguro_consolidado["coberturas"].append(c)

                # Pausa para refrescar el límite OTPM de Groq
                time.sleep(12)

            datos_consolidados.append(seguro_consolidado)
            print(f"✅ Guardado: {seguro_consolidado['nombre_seguro']} con {len(seguro_consolidado['coberturas'])} coberturas.")

    salida_json = Path("matriz_coberturas.json")
    with open(salida_json, "w", encoding="utf-8") as f:
        json.dump(datos_consolidados, f, ensure_ascii=False, indent=2)

    print(f"\n🎉 ¡Finalizado con éxito! Se guardó en '{salida_json.name}' con {len(datos_consolidados)} seguros.")

if __name__ == "__main__":
    procesar_archivos()
